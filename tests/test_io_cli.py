import json

import openpyxl
import pandas as pd
import pytest

from data_cleaning_agent.cli import main
from data_cleaning_agent.contracts import StageResult
from data_cleaning_agent.io import export_result, read_table


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_round_trip_preserves_arabic_leading_zeros_and_literal_formula(tmp_path, extension):
    df = pd.DataFrame(
        {"الاسم": ["مثال", "اختبار"], "id": ["00123", "00045"], "feedback": ["=1+1", "NA"]}
    )
    data, report = export_result(StageResult(df, "text"), tmp_path / f"demo.{extension}")
    restored = read_table(data)
    pd.testing.assert_frame_equal(restored, df)
    assert json.loads(report.read_text())["pipeline_validated"] is False
    if extension == "xlsx":
        workbook = openpyxl.load_workbook(data)
        assert workbook["text"]["C2"].data_type == "s"
        workbook.close()


def test_no_overwrite_even_when_report_already_exists(tmp_path):
    df = pd.DataFrame({"value": ["x"]})
    output = tmp_path / "demo.csv"
    report = tmp_path / "demo.text.report.json"
    report.write_text("existing")
    with pytest.raises(FileExistsError):
        export_result(StageResult(df, "text"), output)
    assert not (tmp_path / "demo.text.csv").exists()
    assert report.read_text() == "existing"


def test_prevent_source_overwrite(tmp_path):
    source = tmp_path / "demo.text.csv"
    source.write_text("original")
    with pytest.raises(ValueError, match="replace the source"):
        export_result(StageResult(pd.DataFrame(), "text"), source, source=source)
    assert source.read_text() == "original"


def test_sheet_header_selection_keeps_checkbox_rows_and_numeric_ids(tmp_path):
    source = tmp_path / "source.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "empty"
    sheet = workbook.create_sheet("responses")
    sheet.append(["Column1", "Column2"])
    sheet.append(["id", "attendance"])
    sheet.append([12345, True])
    sheet.append([None, False])
    workbook.save(source)
    workbook.close()
    df = read_table(source, sheet="responses", header_row=2)
    assert list(df.columns) == ["id", "attendance"]
    assert len(df) == 2 and df.at[1, "attendance"] is False
    assert df.at[0, "id"] == 12345  # No invented leading zeros.
    assert df.attrs["source"]["first_data_row"] == 3
    with pytest.raises(ValueError, match="empty"):
        read_table(source)


def test_cli_text_and_repeated_output(tmp_path):
    output = tmp_path / "result.csv"
    args = [
        "stage",
        "text",
        "--input",
        "examples/synthetic.csv",
        "--column",
        "feedback",
        "--output",
        str(output),
    ]
    assert main(args) == 0
    assert (tmp_path / "result.text.csv").exists()
    assert main(args) == 2


def test_cli_private_model_input_blocked_before_read(tmp_path, capsys):
    assert (
        main(
            [
                "stage",
                "schema",
                "--input",
                "does-not-exist.xlsx",
                "--output",
                str(tmp_path / "out.csv"),
            ]
        )
        == 2
    )
    assert "privacy" in capsys.readouterr().out
    assert not list(tmp_path.iterdir())


def test_cli_missing_pipeline_and_status(capsys):
    assert main(["pipeline"]) == 2
    assert "validation" in capsys.readouterr().out
    assert main(["status"]) == 0


def test_cli_rejects_malformed_prepared_input_without_echo(tmp_path, capsys):
    source = tmp_path / "bad.json"
    source.write_text(json.dumps({"provenance": "raw_private_value"}))
    assert (
        main(
            [
                "stage",
                "schema",
                "--sanitized-input",
                str(source),
                "--output",
                str(tmp_path / "out.csv"),
            ]
        )
        == 2
    )
    assert "raw_private_value" not in capsys.readouterr().out
